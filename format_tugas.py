from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

INPUT_FILE = "tugas.docx"
OUTPUT_FILE = "tugas_formatted.docx"

document = Document(INPUT_FILE)

# ==========================
# Margin
# ==========================
for section in document.sections:
    section.top_margin = Cm(2.25)
    section.bottom_margin = Cm(2.25)
    section.left_margin = Cm(2.25)
    section.right_margin = Cm(2.25)

# ==========================
# Tambah Header Atas
# ==========================

first_para = document.paragraphs[0]

header_para = first_para.insert_paragraph_before(
    "22647005 Maxibiano Soares Horaçio"
)

header_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

for run in header_para.runs:
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(11)

# ==========================
# Format Semua Paragraph
# ==========================

for para in document.paragraphs:

    para.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY

    para.paragraph_format.line_spacing = 1
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(0)

    for run in para.runs:
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)

document.save(OUTPUT_FILE)

print("Selesai!")
print("Output:", OUTPUT_FILE)